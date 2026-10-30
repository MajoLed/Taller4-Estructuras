#!/usr/bin/env python3
"""
Genera la hoja de cálculo del Taller 4 (resultados/Taller4-resultados.xlsx) a partir de los CSV que
escribe `./gradlew experimentos`: una hoja por CSV con los datos tabulados y gráficas log-log con
línea de tendencia POTENCIAL mostrando ecuación y R² (igual que en Excel/LibreOffice).

Requiere openpyxl (pip install openpyxl). Uso:
    python3 scripts/generar_hoja.py [directorio_resultados]

Las gráficas de ajuste usan solo la serie geométrica N = 6 250 … 200 000 (razón constante 2), como
pide el enunciado; el texto completo queda en las tablas.
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.chart.trendline import Trendline
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DIR = sys.argv[1] if len(sys.argv) > 1 else "resultados"
SALIDA = os.path.join(DIR, "Taller4-resultados.xlsx")
TAMANOS = [6250, 12500, 25000, 50000, 100000, 200000]
NEGRITA = Font(bold=True)
ENCABEZADO = PatternFill("solid", fgColor="DDEBF7")


def numero(v):
    if v in ("true", "false"):
        return v == "true"
    try:
        return int(v)
    except ValueError:
        try:
            return float(v)
        except ValueError:
            return v


def leer(nombre):
    ruta = os.path.join(DIR, nombre)
    if not os.path.exists(ruta):
        return None, []
    with open(ruta, encoding="utf-8") as f:
        filas = list(csv.reader(f))
    return filas[0], [[numero(v) for v in fila] for fila in filas[1:]]


def volcar(ws, encabezado, filas, fila0=1, col0=1):
    for j, h in enumerate(encabezado):
        c = ws.cell(row=fila0, column=col0 + j, value=h)
        c.font = NEGRITA
        c.fill = ENCABEZADO
        c.alignment = Alignment(wrap_text=True)
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila):
            ws.cell(row=fila0 + 1 + i, column=col0 + j, value=v)
    for j in range(len(encabezado)):
        letra = get_column_letter(col0 + j)
        ws.column_dimensions[letra].width = max(ws.column_dimensions[letra].width or 0, 12)


def hoja_csv(wb, titulo, nombre):
    enc, filas = leer(nombre)
    if enc is None:
        return None, None, []
    ws = wb.create_sheet(titulo)
    volcar(ws, enc, filas)
    ws.freeze_panes = "A2"
    return ws, enc, filas


def grafica_loglog(titulo, xt, yt, logx=True, logy=True):
    ch = ScatterChart()
    ch.title = titulo
    ch.style = 13
    ch.x_axis.title = xt
    ch.y_axis.title = yt
    ch.height, ch.width = 9, 16
    if logx:
        ch.x_axis.scaling.logBase = 10
    if logy:
        ch.y_axis.scaling.logBase = 10
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    return ch


def serie(ws, col_x, col_y, f1, f2, nombre, tendencia=True):
    xs = Reference(ws, min_col=col_x, min_row=f1, max_row=f2)
    ys = Reference(ws, min_col=col_y, min_row=f1, max_row=f2)
    s = Series(ys, xs, title=nombre)
    s.marker.symbol = "circle"
    s.marker.size = 6
    s.graphicalProperties.line.noFill = True
    if tendencia:
        s.trendline = Trendline(trendlineType="power", dispRSqr=True, dispEq=True)
    return s


def bloque(ws, fila0, col0, titulo, encabezado, filas):
    """Escribe un bloque auxiliar de datos (para graficar) y devuelve (fila_inicio_datos, fila_fin)."""
    ws.cell(row=fila0, column=col0, value=titulo).font = NEGRITA
    volcar(ws, encabezado, filas, fila0 + 1, col0)
    return fila0 + 2, fila0 + 1 + len(filas)


def main():
    wb = Workbook()
    resumen = wb.active
    resumen.title = "Resumen"
    texto = [
        ("Taller práctico 4 — Colas de prioridad y árboles binarios de búsqueda", True),
        ("Equipo: <completar nombres e ID>", False),
        ("Combinación asignada (./gradlew seleccion --args=\"<ids>\"): <completar hash y combinación>", False),
        ("Texto: pg2000.txt (Project Gutenberg #2000) — fecha de descarga: <completar>", False),
        ("", False),
        ("Valores de verificación obtenidos (las 9 combinaciones)", True),
        ("N = 383 637   M = 22 944   Q = 295", False),
        ("Top-10: que (20769), de (18410), y (18272), la (10492), a (9876), en (8285), el (8267), no (6346), los (4769), se (4752)", False),
        ("dul (R = 12): dulcinea (286), dulce (26), dulces (5), dulcísima (4), dulcineas (2)", False),
        ("san (R = 53): sancho (2174), sansón (77), sangre (65), santa (49), san (32)", False),
        ("rank(\"quijote\") = 17 881   min() = a   max() = útiles", False),
        ("", False),
        ("Hojas: una por CSV de resultados/ y hojas de gráficas con línea de tendencia potencial (ecuación y R²).", False),
        ("Los ajustes calculados por el programa (a, b, R²) están en la hoja 'ajustes' (archivos ajustes_<inciso>.csv).", False),
        ("Tiempos en milisegundos; *_sd = desviación estándar muestral de las repeticiones.", False),
    ]
    for i, (t, b) in enumerate(texto, start=1):
        resumen.cell(row=i, column=1, value=t).font = Font(bold=b)
    resumen.column_dimensions["A"].width = 120

    # ------------------------------------------------------------ 6a
    ws, enc, filas = hoja_csv(wb, "6a_heaps", "6a_heaps.csv")
    if ws:
        ch = grafica_loglog("Ley de Heaps: M vs N (serie geométrica)", "N (palabras)", "M (vocabulario)")
        f2 = 1 + sum(1 for f in filas if f[enc.index("en_ajuste")] is True)
        ch.series.append(serie(ws, 1, 2, 2, f2, "M"))
        ws.add_chart(ch, "H2")
    ws, enc, filas = hoja_csv(wb, "6a_zipf", "6a_zipf.csv")
    if ws:
        ch = grafica_loglog("Ley de Zipf: frecuencia vs rango (top-1000)", "rango", "frecuencia")
        ch.series.append(serie(ws, 1, 3, 2, 1 + len(filas), "frecuencia"))
        ws.add_chart(ch, "H2")

    # ------------------------------------------------------------ 6b
    ws, enc, filas = hoja_csv(wb, "6b_escalamiento", "6b_escalamiento.csv")
    if ws:
        g = wb.create_sheet("6b_graficas")
        i_t, i_c, i_n = enc.index("tabla"), enc.index("cola"), enc.index("N")
        combos = []
        for f in filas:
            if (f[i_t], f[i_c]) not in combos:
                combos.append((f[i_t], f[i_c]))
        fila = 1
        for (t, c) in combos:
            datos = [f for f in filas if f[i_t] == t and f[i_c] == c and f[i_n] in TAMANOS]
            cols = ["N", "tA_ms", "tB_ms", "tC_ms", "A_compares", "A_moves", "A_rotations", "A_flips"]
            bl = [[f[enc.index(k)] for k in cols] for f in datos]
            f1, f2 = bloque(g, fila, 1, f"{t} + {c}", cols, bl)
            ch = grafica_loglog(f"{t} + {c}: tiempo por fase", "N", "tiempo (ms)")
            for j, nombre in ((2, "fase A"), (3, "fase B"), (4, "fase C")):
                ch.series.append(serie(g, 1, j, f1, f2, nombre))
            g.add_chart(ch, f"J{fila}")
            ch2 = grafica_loglog(f"{t} + {c}: contadores de la fase A", "N", "operaciones")
            for j, nombre in ((5, "compares"), (6, "moves"), (7, "rotations"), (8, "flips")):
                if all(isinstance(r[j - 1], (int, float)) and r[j - 1] > 0 for r in bl):
                    ch2.series.append(serie(g, 1, j, f1, f2, nombre))
            g.add_chart(ch2, f"T{fila}")
            fila += 20

    # ------------------------------------------------------------ 6c
    ws, enc, filas = hoja_csv(wb, "6c_forma", "6c_forma.csv")
    if ws:
        g = wb.create_sheet("6c_graficas")
        i_n, i_a = enc.index("N"), enc.index("arbol")
        ns = []
        for f in filas:
            if f[i_n] not in ns:
                ns.append(f[i_n])
        def val(n, arbol, col):
            return next(f[enc.index(col)] for f in filas if f[i_n] == n and f[i_a] == arbol)
        cols = ["N", "h BST", "dBar BST", "dBarF BST", "h RB", "dBar RB", "dBarF RB", "lg M", "2 ln M", "2 lg M"]
        bl = [[n, val(n, "BST", "h"), val(n, "BST", "dBar"), val(n, "BST", "dBarF"),
               val(n, "RedBlackBST", "h"), val(n, "RedBlackBST", "dBar"), val(n, "RedBlackBST", "dBarF"),
               val(n, "BST", "lgM"), val(n, "BST", "dos_lnM"), val(n, "BST", "dos_lgM")] for n in ns]
        f1, f2 = bloque(g, 1, 1, "Forma de los árboles (incluye el texto completo)", cols, bl)
        ch = grafica_loglog("h y d̄ frente a lg M, 2 ln M y 2 lg M", "N", "enlaces", logy=False)
        for j in range(2, 11):
            if cols[j - 1].startswith("dBarF"):
                continue
            s = serie(g, 1, j, f1, f2, cols[j - 1], tendencia=False)
            if j >= 8:
                s.marker.symbol = "none"
                s.graphicalProperties.line.noFill = False
                s.graphicalProperties.line.dashStyle = "dash"
            else:
                s.graphicalProperties.line.noFill = False
            ch.series.append(s)
        g.add_chart(ch, "L2")
        ch2 = grafica_loglog("d̄ frente a d̄_f", "N", "profundidad", logy=False)
        for j in (3, 4, 6, 7):
            s = serie(g, 1, j, f1, f2, cols[j - 1], tendencia=False)
            s.graphicalProperties.line.noFill = False
            ch2.series.append(s)
        g.add_chart(ch2, "L22")

    # ------------------------------------------------------------ 6d
    ws, enc, filas = hoja_csv(wb, "6d_orden", "6d_orden.csv")
    if ws:
        g = wb.create_sheet("6d_graficas")
        i_t, i_e, i_s = enc.index("tabla"), enc.index("escenario"), enc.index("semilla")
        tablas = []
        for f in filas:
            if f[i_t] not in tablas:
                tablas.append(f[i_t])
        escenarios = []
        for f in filas:
            nombre = f[i_e] if f[i_s] == 0 else f"{f[i_e]} s={f[i_s]}"
            if nombre not in escenarios:
                escenarios.append(nombre)
        def fila_de(t, esc):
            for f in filas:
                nombre = f[i_e] if f[i_s] == 0 else f"{f[i_e]} s={f[i_s]}"
                if f[i_t] == t and nombre == esc:
                    return f
            return None
        fila = 1
        for metrica, titulo in (("compares", "comparaciones de la construcción"),
                                ("compares_busqueda", "comparaciones de la carga de búsqueda (N get)"),
                                ("tBusqueda_ms", "tiempo de la carga de búsqueda (ms)")):
            bl = [[esc] + [fila_de(t, esc)[enc.index(metrica)] if fila_de(t, esc) else None for t in tablas] for esc in escenarios]
            f1, f2 = bloque(g, fila, 1, titulo, ["escenario"] + tablas, bl)
            ch = BarChart()
            ch.type = "col"
            ch.title = titulo
            ch.y_axis.scaling.logBase = 10
            ch.x_axis.delete = False
            ch.y_axis.delete = False
            ch.height, ch.width = 9, 20
            ch.add_data(Reference(g, min_col=2, max_col=1 + len(tablas), min_row=f1 - 1, max_row=f2), titles_from_data=True)
            ch.set_categories(Reference(g, min_col=1, min_row=f1, max_row=f2))
            g.add_chart(ch, f"F{fila}")
            fila += 20

    # ------------------------------------------------------------ 6e
    ws, enc, filas = hoja_csv(wb, "6e_efectoK", "6e_efectoK.csv")
    if ws:
        g = wb.create_sheet("6e_graficas")
        i_c = enc.index("cola")
        colas = []
        for f in filas:
            if f[i_c] not in colas:
                colas.append(f[i_c])
        fila = 1
        ch_l = grafica_loglog("Fase B: less vs K", "K", "less")
        ch_x = grafica_loglog("Fase B: exch vs K", "K", "exch")
        ch_t = grafica_loglog("Fase B: tiempo vs K", "K", "tiempo (ms)")
        for c in colas:
            bl = [[f[enc.index("K")], f[enc.index("less")], f[enc.index("exch")], f[enc.index("t_ms")]] for f in filas if f[i_c] == c]
            f1, f2 = bloque(g, fila, 1, c, ["K", "less", "exch", "t_ms"], bl)
            ch_l.series.append(serie(g, 1, 2, f1, f2, c))
            if all(r[2] > 0 for r in bl):
                ch_x.series.append(serie(g, 1, 3, f1, f2, c))
            ch_t.series.append(serie(g, 1, 4, f1, f2, c))
            fila += len(bl) + 3
        g.add_chart(ch_l, "G1")
        g.add_chart(ch_x, "G20")
        g.add_chart(ch_t, "G39")
    ws, enc, filas = hoja_csv(wb, "6e_E1_vs_E2", "6e_E1_vs_E2.csv")
    if ws:
        n = 1 + len(filas)
        ch = grafica_loglog("E1 (MinPQ acotada) vs E2 (heapify + K delMax): less", "K", "less")
        ch.series.append(serie(ws, 1, enc.index("E1_less") + 1, 2, n, "E1", tendencia=False))
        ch.series.append(serie(ws, 1, enc.index("E2_less") + 1, 2, n, "E2", tendencia=False))
        for s in ch.series:
            s.graphicalProperties.line.noFill = False
        ws.add_chart(ch, "P2")
        ch2 = grafica_loglog("E1 vs E2: tiempo", "K", "tiempo (ms)")
        ch2.series.append(serie(ws, 1, enc.index("E1_ms") + 1, 2, n, "E1", tendencia=False))
        ch2.series.append(serie(ws, 1, enc.index("E2_ms") + 1, 2, n, "E2", tendencia=False))
        for s in ch2.series:
            s.graphicalProperties.line.noFill = False
        ws.add_chart(ch2, "P22")

    # ------------------------------------------------------------ 6f y ajustes
    hoja_csv(wb, "6f_referencia", "6f_referencia.csv")
    ajustes, enc = [], None
    for nombre in sorted(os.listdir(DIR)):
        if nombre.startswith("ajustes_") and nombre.endswith(".csv"):
            e, filas = leer(nombre)
            enc = e
            ajustes += filas
    if enc:
        ws = wb.create_sheet("ajustes")
        volcar(ws, enc, ajustes)
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["C"].width = 26

    wb.save(SALIDA)
    print("Hoja generada:", SALIDA)


if __name__ == "__main__":
    main()
